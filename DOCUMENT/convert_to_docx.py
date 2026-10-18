import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=4, val='single', color='003366')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "auto")}"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def create_styled_document(md_path, docx_path):
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def process_table(lines):
        # Extract rows
        rows_data = []
        for line in lines:
            line_str = line.strip()
            if not line_str.startswith('|'):
                continue
            # Check if separator
            if re.match(r'^\|[\s\-:|]+\|$', line_str):
                continue
            cells = [c.strip() for c in line_str.split('|')[1:-1]]
            rows_data.append(cells)
            
        if not rows_data:
            return

        col_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Determine if it's an image placeholder table or standard table
        is_image_box = False
        if len(rows_data) > 0 and any("CHÈN ẢNH" in c for c in rows_data[0]):
            is_image_box = True

        for r_idx, row in enumerate(rows_data):
            # Pad if needed
            while len(row) < col_count:
                row.append("")
                
            for c_idx, cell_value in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # Style cell
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                
                # Format cell text
                # Check markdown formatting: **bold**, `code`, <br>
                parts = re.split(r'(<br\s*/?>)', cell_value)
                for part in parts:
                    if re.match(r'<br\s*/?>', part):
                        p = cell.add_paragraph()
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        continue
                    
                    # Split for bold and inline code
                    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', part)
                    for tok in tokens:
                        if tok.startswith('**') and tok.endswith('**'):
                            run = p.add_run(tok[2:-2])
                            run.bold = True
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(10.5)
                        elif tok.startswith('`') and tok.endswith('`'):
                            run = p.add_run(tok[1:-1])
                            run.font.name = 'Consolas'
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
                        else:
                            run = p.add_run(tok)
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(10.5)

                # Shading & borders
                if is_image_box:
                    set_cell_background(cell, "F4F6F9")
                    set_cell_border(cell, 
                                    top=dict(val='single', sz=6, color='3498DB'),
                                    bottom=dict(val='single', sz=6, color='3498DB'),
                                    left=dict(val='single', sz=6, color='3498DB'),
                                    right=dict(val='single', sz=6, color='3498DB'))
                elif r_idx == 0:
                    # Header row
                    set_cell_background(cell, "1F4E79")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.bold = True
                    set_cell_border(cell, 
                                    top=dict(val='single', sz=8, color='14385C'),
                                    bottom=dict(val='single', sz=8, color='14385C'),
                                    left=dict(val='single', sz=4, color='D3D3D3'),
                                    right=dict(val='single', sz=4, color='D3D3D3'))
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "FFFFFF")
                    else:
                        set_cell_background(cell, "F9FBFD")
                    set_cell_border(cell, 
                                    top=dict(val='single', sz=4, color='E0E0E0'),
                                    bottom=dict(val='single', sz=4, color='E0E0E0'),
                                    left=dict(val='single', sz=4, color='E0E0E0'),
                                    right=dict(val='single', sz=4, color='E0E0E0'))

        # Set column widths nicely
        if col_count == 2:
            table.columns[0].width = Inches(1.8)
            table.columns[1].width = Inches(4.7)
        elif col_count == 5:
            table.columns[0].width = Inches(1.2)
            table.columns[1].width = Inches(1.0)
            table.columns[2].width = Inches(1.8)
            table.columns[3].width = Inches(1.5)
            table.columns[4].width = Inches(1.0)
            
        doc.add_paragraph() # Spacing after table

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check table
        if line.strip().startswith('|'):
            table_lines.append(line)
            i += 1
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].rstrip('\r\n'))
                i += 1
            process_table(table_lines)
            table_lines = []
            continue
            
        # Headers
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(line[2:])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[3:])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[5:])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        elif line.startswith('##### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[6:])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
        elif line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('_______________________________________________________________________________')
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(8)
        elif line.startswith('*Hình ') and line.endswith('*'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[1:-1])
            run.italic = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif line.startswith('```'):
            # Code block
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\r\n'))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        elif line.strip().startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            content = line.strip()[2:]
            tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', content)
            for tok in tokens:
                if tok.startswith('**') and tok.endswith('**'):
                    run = p.add_run(tok[2:-2])
                    run.bold = True
                elif tok.startswith('`') and tok.endswith('`'):
                    run = p.add_run(tok[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
                else:
                    run = p.add_run(tok)
        elif line.strip() != "":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', line)
            for tok in tokens:
                if tok.startswith('**') and tok.endswith('**'):
                    run = p.add_run(tok[2:-2])
                    run.bold = True
                elif tok.startswith('*') and tok.endswith('*'):
                    run = p.add_run(tok[1:-1])
                    run.italic = True
                elif tok.startswith('`') and tok.endswith('`'):
                    run = p.add_run(tok[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
                else:
                    run = p.add_run(tok)
        i += 1

    doc.save(docx_path)
    print(f"Document saved successfully: {docx_path}")

if __name__ == "__main__":
    md_file = r"c:\Users\HP\Documents\the4BookStore\DOCUMENT\kiem_thu_he_thong_hoan_chinh.md"
    docx_file = r"c:\Users\HP\Documents\the4BookStore\DOCUMENT\Bao_Cao_Kiem_Thu_He_Thong_The4BookStore.docx"
    create_styled_document(md_file, docx_file)
